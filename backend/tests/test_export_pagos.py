"""Reporte de pagos descargable en Excel (GET /pagos/export)."""
import io
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.deps import get_db
from app.core.security import create_access_token, hash_password
from app.main import app
from app.models.pago import Pago
from app.models.profesor import Profesor
from app.models.usuario import Usuario


@pytest.fixture
def client(db):
    app.dependency_overrides[get_db] = lambda: db
    yield TestClient(app)
    app.dependency_overrides.clear()


def _headers(db, username, rol):
    db.add(Usuario(username=username, nombre=username, password_hash=hash_password("x"),
                   rol=rol, activo=True))
    db.flush()
    return {"Authorization": f"Bearer {create_access_token(username, rol)}"}


def test_revisor_descarga_excel_de_pagos(client, db):
    headers = _headers(db, "rev_pag", "revisor")
    p = Profesor(rfc="PAG010101AA1", nombre="Profe Pagado", correo="pag@example.com",
                 regimen_fiscal="612", activo=True)
    db.add(p)
    db.flush()
    db.add(Pago(profesor_id=p.id, mes=6, anio=2026, pagada=True,
                fecha_pago=date(2026, 6, 20), metodo_pago="Transferencia",
                registrado_por="rev_pag"))
    db.flush()

    r = client.get("/api/v1/pagos/export?mes=6&anio=2026", headers=headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Verifac_pagos_2026-06.xlsx" in r.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    encabezados = [c.value for c in ws[1]]
    assert "Pagada" in encabezados and "Fecha de pago" in encabezados
    filas = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    fila = filas["Profe Pagado"]
    assert fila[encabezados.index("Pagada")].value == "Sí"
    assert fila[encabezados.index("Método de pago")].value == "Transferencia"


def test_consulta_no_descarga_pagos(client, db):
    headers = _headers(db, "lector_pag", "consulta")
    assert client.get("/api/v1/pagos/export?mes=6&anio=2026", headers=headers).status_code == 403
